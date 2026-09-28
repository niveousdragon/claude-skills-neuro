"""Fill an MSU "заключение о возможности открытого опубликования" blank.

The faculty memo allows the author to change ONLY the yellow-highlighted text
of the blank; everything else is filled in by the commission. So this script
touches nothing but the highlighted runs, and removes the highlight, which the
memo requires before printing.

Each page of a blank has three highlighted places, in this order:
  1. the header line under "о возможности открытого опубликования"
     (kind of material, authors, title, and where it is going);
  2. "провела экспертизу материалов <...>";
  3. "Сведения, содержащиеся в рассматриваемых материалах <...>".
Places 2 and 3 carry the same text, so they are given once.

Usage:
  python fill_clearance.py BLANK.docx --list
  python fill_clearance.py BLANK.docx OUT.docx --page "HEADER" "MATERIAL" [--page ...]

Text between *asterisks* is set in italics (the blanks italicise the title).
A space after author initials becomes a non-breaking one, as in the samples.
If MS Word is available, the page count is checked and a PDF is written next
to OUT for a visual check; the blank must print on exactly one page per form.
"""
import argparse
import copy
import os
import re
import shutil
import sys
import tempfile

import docx
from docx.enum.text import WD_COLOR_INDEX

PLACES_PER_PAGE = 3


def highlighted_groups(document):
    """Runs of consecutive highlighted text, grouped per paragraph, in order."""
    groups = []
    for par in document.paragraphs:
        current = []
        for run in par.runs:
            if run.font.highlight_color not in (None, WD_COLOR_INDEX.AUTO):
                current.append(run)
            elif current:
                groups.append(current)
                current = []
        if current:
            groups.append(current)
    return groups


def nbsp_after_initials(text):
    # "Н.А. Поспелова" -> "Н.А.\u00a0Поспелова"; matches the samples in the blanks.
    return re.sub(r"((?:\b[А-ЯЁA-Z]\.){1,2}) +(?=[А-ЯЁA-Z])", "\\1\u00a0", text)


def replace_group(group, text):
    template = group[0]._r
    parent = template.getparent()
    anchor = template
    for i, chunk in enumerate(re.split(r"\*", nbsp_after_initials(text))):
        if not chunk:
            continue
        new = copy.deepcopy(template)
        parent.insert(parent.index(anchor) + 1, new)
        anchor = new
        run = docx.text.run.Run(new, group[0]._parent)
        run.text = chunk
        run.italic = True if i % 2 else None
        run.font.highlight_color = None
    for run in group:
        run._r.getparent().remove(run._r)


def word_check(out_path):
    """Return the page count from MS Word and write a PDF, or None without Word."""
    try:
        import win32com.client
    except ImportError:
        return None
    tmp = tempfile.mkdtemp()
    # Word fails to open paths with non-ASCII characters on some setups.
    src = os.path.join(tmp, "blank.docx")
    shutil.copy(out_path, src)
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(src, ReadOnly=True)
        pages = doc.ComputeStatistics(2)  # wdStatisticPages
        doc.SaveAs2(os.path.join(tmp, "blank.pdf"), FileFormat=17)  # wdFormatPDF
        doc.Close(False)
    finally:
        word.Quit()
    shutil.copy(os.path.join(tmp, "blank.pdf"), os.path.splitext(out_path)[0] + ".pdf")
    return pages


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("blank")
    ap.add_argument("out", nargs="?")
    ap.add_argument("--list", action="store_true", help="show the highlighted places and exit")
    ap.add_argument("--page", nargs=2, action="append", metavar=("HEADER", "MATERIAL"), default=[])
    args = ap.parse_args()
    # The sample text is Cyrillic; a cp1251 console or a UTF-8 pipe must both read it.
    sys.stdout.reconfigure(encoding="utf-8")

    document = docx.Document(args.blank)
    groups = highlighted_groups(document)

    if args.list or not args.out:
        for i, g in enumerate(groups):
            print(f"[{i // PLACES_PER_PAGE + 1}.{i % PLACES_PER_PAGE + 1}] {''.join(r.text for r in g)}")
        return 0

    if len(groups) != PLACES_PER_PAGE * len(args.page):
        print(f"[FAILED] the blank has {len(groups)} highlighted places, expected "
              f"{PLACES_PER_PAGE} per page x {len(args.page)} --page. Either the number of --page "
              "does not match the blank, or the blank's layout changed: run --list and compare "
              "with the memo before filling it by hand.")
        return 1

    for n, (header, material) in enumerate(args.page):
        first = n * PLACES_PER_PAGE
        replace_group(groups[first], header)
        replace_group(groups[first + 1], material)
        replace_group(groups[first + 2], material)
    document.save(args.out)

    left = highlighted_groups(docx.Document(args.out))
    if left:
        print(f"[FAILED] highlight left in {len(left)} places")
        return 1
    print(f"[OK] filled {len(groups)} places, highlight removed: {args.out}")

    pages = word_check(args.out)
    expected = len(args.page)
    if pages is None:
        print("[WARN] MS Word not available: page count NOT checked. Open the file and make sure "
              f"each form fits on one page ({expected} page(s) in total).")
    elif pages != expected:
        print(f"[FAILED] {pages} pages, expected {expected}: a form spilled onto the next page. "
              "Show the PDF to the user; do not touch fonts, margins or text outside the filled places.")
        return 1
    else:
        print(f"[OK] {pages} page(s), one per form; PDF for a visual check: "
              f"{os.path.splitext(args.out)[0]}.pdf")
    return 0


if __name__ == "__main__":
    sys.exit(main())
